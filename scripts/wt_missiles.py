"""从 E:\\WT\\WT.xlsx 的 "AAM (ARH)" 表提取权威 (ΔV, BC)，输出给交叉图与工作流使用。

口径（与 AGENTS.md 一致，逐格照抄表内公式）：
    Isp1 = t1·F1/(9.81·m1)                            row 13
    ΔV1  = Isp1·9.81·ln(m0/(m0-m1))                   row 14
    Isp2 = t2·F2/(9.81·m2)                            row 18
    ΔV2  = Isp2·9.81·ln((m0-m1)/(m0-m1-m2))           row 19
    Total ΔV = ΔV1 + ΔV2（无二级取 ΔV1）               row 20
    BC   = (m0-m1-m2)/(k·(d/2000)²·π)                 row 87（m2 为 "-" 时取 0）

用法：
    python scripts/wt_missiles.py            # 打印并写 results/wt_missiles.json
    from wt_missiles import load_wt_missiles # 返回 [{name, dv, bc, ...}]
"""

from __future__ import annotations

import json
import math
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(r"E:\WT\WT.xlsx")
SHEET = "AAM (ARH)"
OUT = ROOT / "results" / "wt_missiles.json"

ROW = {"name": 1, "m0": 3, "m1": 4, "m2": 5, "d_mm": 6,
       "f1": 10, "t1": 11, "f2": 16, "t2": 17, "dv": 20, "k": 86, "bc": 87}

# 交叉图里要用的代表列（表里有重复列时显式指定）
PREFERRED = [
    "AIM-120A (2.58.0.27)", "AIM-120C-5 (2.58.0.27)", "AIM-120D (2.58.0.27)",
    "AAM-4", "Derby", "R-Darter",
    "MICA EM (new)", "PL-12", "PL-12A", "R-77", "R-77-1",
]
# 图上标签 → 表内列名（重复列导致的歧义在这里消解）
FIGURE_LABELS = {
    "120A": "AIM-120A (2.58.0.27)",
    "120C": "AIM-120C-5 (2.58.0.27)",
    "120D": "AIM-120D (2.58.0.27)",
    "AAM-4": "AAM-4",
    "Derby": "Derby",
    "R-Darter": "R-Darter",
    "MICA": "MICA EM (new)",
    "PL-12": "PL-12",
    "PL-12A": "PL-12A",
    "R-77": "R-77",
    "R-77-1": "R-77-1",
}


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if s in ("", "-", "n/a"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _compute(m0, m1, m2, f1, t1, f2, t2, k, d_mm):
    if not (m0 and m1 and k and d_mm):
        return None
    isp1 = (t1 * f1 / (9.81 * m1)) if (t1 and f1) else None
    dv1 = (isp1 * 9.81 * math.log(m0 / (m0 - m1))) if isp1 else None
    isp2 = dv2 = None
    if m2 and f2 and t2:
        isp2 = t2 * f2 / (9.81 * m2)
        dv2 = isp2 * 9.81 * math.log((m0 - m1) / (m0 - m1 - m2))
    dv = None if (dv1 is None and dv2 is None) else (dv1 or 0.0) + (dv2 or 0.0)
    bc = (m0 - m1 - (m2 or 0.0)) / (k * (d_mm / 2000.0) ** 2 * math.pi)
    return {"isp1": isp1, "dv1": dv1, "isp2": isp2, "dv2": dv2, "dv": dv, "bc": bc}


def read_sheet(xlsx: Path = XLSX, sheet: str = SHEET) -> dict:
    """读取整张表，返回 {列名: {参数..., computed...}}。"""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[sheet]
    out = {}
    for c in range(2, ws.max_column + 1):
        name = ws.cell(row=ROW["name"], column=c).value
        if not name:
            continue
        rec = {"col": openpyxl.utils.get_column_letter(c), "name": str(name).strip()}
        for key, r in ROW.items():
            if key == "name":
                continue
            rec[key] = ws.cell(row=r, column=c).value
        rec["computed"] = _compute(
            _num(rec["m0"]), _num(rec["m1"]), _num(rec["m2"]),
            _num(rec["f1"]), _num(rec["t1"]), _num(rec["f2"]), _num(rec["t2"]),
            _num(rec["k"]), _num(rec["d_mm"]),
        )
        rec["dv_cached"] = _num(rec["dv"])
        rec["bc_cached"] = _num(rec["bc"])
        out[rec["name"]] = rec
    return out


def load_wt_missiles(preferred=None) -> list:
    """返回交叉图用的弹列表（表内权威值）。"""
    table = read_sheet()
    names = preferred if preferred is not None else PREFERRED
    rows = []
    for nm in names:
        rec = table.get(nm)
        if not rec or not rec["computed"]:
            continue
        comp = rec["computed"]
        if comp["dv"] is None or comp["bc"] is None:
            continue
        rows.append({
            "name": nm,
            "dv": float(comp["dv"]),
            "bc": float(comp["bc"]),
            "col": rec["col"],
            "m0": _num(rec["m0"]), "m1": _num(rec["m1"]), "m2": _num(rec["m2"]),
            "d_mm": _num(rec["d_mm"]), "k": _num(rec["k"]),
            "f1": _num(rec["f1"]), "t1": _num(rec["t1"]),
            "f2": _num(rec["f2"]), "t2": _num(rec["t2"]),
            "source": "WT.xlsx AAM (ARH)",
        })
    return rows


def load_all_columns() -> list:
    """返回表内全部列的 (ΔV, BC)（用于散点背景/复核）。"""
    table = read_sheet()
    rows = []
    for nm, rec in table.items():
        comp = rec["computed"]
        if not comp or comp["dv"] is None or comp["bc"] is None:
            continue
        rows.append({"name": nm, "col": rec["col"], "dv": float(comp["dv"]), "bc": float(comp["bc"])})
    return rows


def main() -> int:
    rows = load_wt_missiles()
    print(f"{'弹（表内列名）':<28} {'列':>3} {'ΔV':>8} {'BC':>9}  {'二级?':>5}")
    for r in rows:
        two = "有" if r["t2"] else "无"
        print(f"{r['name']:<28} {r['col']:>3} {r['dv']:8.1f} {r['bc']:9.1f}  {two:>5}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"source": str(XLSX), "sheet": SHEET, "missiles": rows},
                              ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n已写出 {OUT}（{len(rows)} 型）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
